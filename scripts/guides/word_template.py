#!/usr/bin/env python3
"""Create the Word INPUT template product people fill in (Danish).

Usage: python3 word_template.py OUT.docx [--example guide.json]
With --example, the template is pre-filled from a guide JSON (used to test the
round trip Word -> guide JSON -> PDF).

Structure the converter (docx_to_guide.py) relies on:
  - Table 1 = fact sheet (label | value rows)
  - "Overskrift 1" / Heading 1 = optional section heading
  - "Overskrift 2" / Heading 2 = one step
  - normal paragraphs under a step = step text; bulleted list = bullets
  - pasted pictures under a step = screenshots, in order
  - paragraph starting "Tip:" or "Vigtigt:" = callout box
Everything in grey italics starting with ">>" is guidance and is ignored.
"""
import json, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Mm

FACTS = [
    ("Titel", "Hvad skal kunden kunne bagefter? Fx: Opdatér firmware på din XLOCK-lås"),
    ("Kort intro", "1-2 sætninger: hvorfor og hvornår har kunden brug for denne vejledning?"),
    ("Produktlinje", "Fx XLOCK"),
    ("Gælder for", "Produktnavn og STROXX-varenummer, ét pr. linje. Fx: Smart Lock ST-2 | 102-112"),
    ("Du skal bruge", "Det kunden skal have klar, ét pr. linje"),
    ("Tid", "Antal minutter, fx 10"),
    ("Niveau", "Let, Mellem eller Avanceret, og hvad det kræver. Fx: Let | Kræver en papirclips"),
    ("Godt at vide", "Den ene ting kunden skal vide før start (valgfri)"),
    ("Når du er færdig", "Hvad har kunden opnået? Fx: Låsen kører nu den nyeste firmware."),
    ("Skrevet af", "Dit navn"),
]


def guide_note(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(">> " + text); r.italic = True; r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x8A, 0x91, 0x99)


def build(out, example=None):
    g = json.loads(Path(example).read_text(encoding="utf-8")) if example else None
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    doc.add_heading("STROXX vejledning", level=0)
    guide_note(doc, "Skriv kun indhold. Design, logo, farver og QR-kode kommer automatisk. "
                    "Send dokumentet til Claude med teksten: \"Lav en STROXX-vejledning af dette\".")

    t = doc.add_table(rows=0, cols=2); t.style = "Table Grid"
    for k, hint in FACTS:
        row = t.add_row().cells; row[0].text = k
        val = hint
        if g:
            val = {
                "Titel": g["title"], "Kort intro": g.get("summary", ""), "Produktlinje": g.get("product_line", ""),
                "Gælder for": "\n".join(f"{a['name']} | {a.get('item_no','')}" for a in g.get("applies_to", [])),
                "Du skal bruge": "\n".join(g.get("you_need", [])), "Tid": g.get("time_num", ""),
                "Niveau": f"{g.get('level','')} | {g.get('level_note','')}", "Godt at vide": g.get("good_to_know", ""),
                "Når du er færdig": g.get("done", ""),
                "Skrevet af": g.get("source", ""),
            }[k]
        row[1].text = val
        row[0].width = Mm(35); row[1].width = Mm(125)

    doc.add_paragraph()
    guide_note(doc, "TRIN: Brug typografien \"Overskrift 2\" til hvert trin (kort, som en handling: \"Slet den defekte lås\"). "
                    "Skriv 1-3 korte sætninger under. Indsæt skærmbilleder direkte under trinnet. "
                    "Start en linje med \"Tip:\" eller \"Vigtigt:\" for at lave en boks. "
                    "Start en linje med \"Du ser nu:\" for at fortælle, hvad kunden skal se, når trinnet er gjort. "
                    "Brug \"Overskrift 1\" hvis vejledningen har flere dele (fx Mulighed 1 / Mulighed 2).")

    if not g:
        doc.add_heading("Trin 1 skrives her som en handling", level=2)
        doc.add_paragraph("Forklar kort hvad kunden skal gøre.")
        guide_note(doc, "Indsæt skærmbillede her (Indsæt > Billeder, eller kopier og sæt ind).")
        doc.add_paragraph("Du ser nu: Hvad kunden ser på skærmen eller låsen, når trinnet er gjort.")
        doc.add_paragraph("Tip: Et godt råd, hvis der er et.")
        doc.add_heading("Trin 2", level=2)
        doc.add_paragraph("…")
    else:
        base = Path(example).parent
        for s in g["sections"]:
            if s.get("heading"): doc.add_heading(s["heading"], level=1)
            if s.get("intro"): doc.add_paragraph(s["intro"])
            for stp in s["steps"]:
                doc.add_heading(stp["title"], level=2)
                if stp.get("text"): doc.add_paragraph(stp["text"])
                for b in stp.get("bullets", []): doc.add_paragraph(b, style="List Bullet")
                for im in stp.get("images", []):
                    doc.add_picture(str(base / im["src"]), width=Mm(55))
                if stp.get("result"): doc.add_paragraph("Du ser nu: " + stp["result"])
                if stp.get("note"):
                    doc.add_paragraph(("Vigtigt: " if stp["note"]["type"] == "warning" else "Tip: ") + stp["note"]["text"])
    doc.save(out); print("wrote", out)


if __name__ == "__main__":
    ex = sys.argv[sys.argv.index("--example") + 1] if "--example" in sys.argv else None
    build(sys.argv[1], ex)
