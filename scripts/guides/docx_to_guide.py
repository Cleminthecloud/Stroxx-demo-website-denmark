#!/usr/bin/env python3
"""Convert a filled-in STROXX Word template into guide JSON + images.

Usage: python3 docx_to_guide.py IN.docx OUTDIR --id english-slug [--help-url stroxx.eu/dk/support/xlock-help]

Writes OUTDIR/guide.json and OUTDIR/img/*. Claude reviews and edits the JSON
(copy edit, item numbers, notes) before rendering with render.py.
"""
import argparse, json, re
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

LABELS = {"titel": "title", "kort intro": "summary", "produktlinje": "product_line",
          "gælder for": "applies_to", "du skal bruge": "you_need", "tid": "time_num", "skrevet af": "source",
          "niveau": "level", "godt at vide": "good_to_know", "når du er færdig": "done"}


def iter_blocks(doc):
    for el in doc.element.body.iterchildren():
        if el.tag == qn("w:p"):
            yield "p", el
        elif el.tag == qn("w:tbl"):
            yield "t", el


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx"); ap.add_argument("out"); ap.add_argument("--id", required=True)
    ap.add_argument("--help-url", default="stroxx.eu/dk/support")
    a = ap.parse_args()
    doc = Document(a.docx)
    out = Path(a.out); (out / "img").mkdir(parents=True, exist_ok=True)

    g = {"id": a.id, "language": "da", "product_line": "", "kind": "Vejledning", "version": "1.0",
         "date": "", "title": "", "summary": "", "applies_to": [], "you_need": [], "time_num": "", "time_unit": "minutter", "level": "Let", "level_note": "",
         "sections": [], "help": {"title": "Virker det ikke?",
         "text": "Find alle STROXX-vejledninger, videoer og manualer samlet ét sted. Scan koden eller gå til",
         "url": a.help_url}}

    # Fact table
    if doc.tables:
        for row in doc.tables[0].rows:
            k = row.cells[0].text.strip().lower(); v = row.cells[1].text.strip()
            key = LABELS.get(k)
            if not key: continue
            if key == "applies_to":
                for line in filter(None, (l.strip() for l in v.splitlines())):
                    name, _, item = line.partition("|")
                    g["applies_to"].append({"name": name.strip(), "item_no": item.strip()})
            elif key == "you_need":
                g["you_need"] = [l.strip() for l in v.splitlines() if l.strip()]
            elif key == "level":
                lv, _, note = v.partition("|"); g["level"] = lv.strip(); g["level_note"] = note.strip()
            elif key == "time_num":
                m = re.search(r"\d+", v); g["time_num"] = m.group(0) if m else v; g["time_unit"] = "minutter"
            else:
                g[key] = v

    section = {"steps": []}; g["sections"].append(section); step = None; n_img = 0
    rels = doc.part.related_parts
    for kind, el in iter_blocks(doc):
        if kind != "p": continue
        from docx.text.paragraph import Paragraph
        p = Paragraph(el, doc)
        style = (p.style.name or "").lower()
        text = p.text.strip()
        if text.startswith(">>"): continue
        blips = el.findall(".//" + qn("a:blip"))
        if style.startswith(("heading 1", "overskrift 1")):
            section = {"heading": text, "steps": []}; g["sections"].append(section); step = None; continue
        if style.startswith(("heading 2", "overskrift 2")):
            step = {"title": text}; section["steps"].append(step); continue
        if style in ("title", "titel") or style.startswith(("heading 0",)): continue
        for b in blips:
            rid = b.get(qn("r:embed")); part = rels.get(rid)
            if part is None or step is None: continue
            n_img += 1; ext = Path(part.partname).suffix or ".png"
            name = f"{a.id}-{n_img:02d}{ext}"; (out / "img" / name).write_bytes(part.blob)
            step.setdefault("images", []).append({"src": f"img/{name}"})
        if not text: continue
        m = re.match(r"^(tip|vigtigt|obs)\s*[:!]\s*(.*)$", text, re.I | re.S)
        if step is None:
            if section.get("heading"): section["intro"] = (section.get("intro", "") + " " + text).strip()
            continue
        r = re.match(r"^du ser nu\s*:\s*(.*)$", text, re.I | re.S)
        if r:
            step["result"] = r.group(1).strip()
        elif m:
            step["note"] = {"type": "tip" if m.group(1).lower() == "tip" else "warning", "text": m.group(2).strip()}
        elif "list" in style:
            step.setdefault("bullets", []).append(text)
        else:
            step["text"] = (step.get("text", "") + "\n\n" + text).strip()

    g["sections"] = [s for s in g["sections"] if s["steps"]]
    (out / "guide.json").write_text(json.dumps(g, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"wrote {out/'guide.json'}: {sum(len(s['steps']) for s in g['sections'])} steps, {n_img} images")


if __name__ == "__main__":
    main()
